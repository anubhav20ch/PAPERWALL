import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report_docx():
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("QuantumShield: Post-Quantum Cryptographic Academic Paper Distribution & Anti-Leak Governance System")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42) # Dark navy

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Comprehensive Technical Architecture, National Exam Leak Resolution, & Quantum Security Governance Report")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run("Author: QuantumShield Security & Architecture Team  |  Date: September 19, 2026  |  Version: 1.0")
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for headings
    def add_h1(text):
        h = doc.add_heading(text, level=1)
        h.style.font.name = "Calibri"
        h.style.font.color.rgb = RGBColor(30, 41, 59)
        h.style.font.size = Pt(16)
        h.style.font.bold = True
        return h

    def add_h2(text):
        h = doc.add_heading(text, level=2)
        h.style.font.name = "Calibri"
        h.style.font.color.rgb = RGBColor(51, 65, 85)
        h.style.font.size = Pt(13)
        h.style.font.bold = True
        return h

    # ABSTRACT
    add_h1("Executive Summary & Abstract")
    p_abs = doc.add_paragraph()
    p_abs.paragraph_format.line_spacing = 1.15
    p_abs.paragraph_format.space_after = Pt(12)
    r_abs = p_abs.add_run(
        "In recent years, high-stakes academic and national competitive examinations worldwide have suffered catastrophic failures "
        "due to paper leakages, unauthorized pre-exam distribution, and insider tampering. Conventional distribution systems rely on "
        "unencrypted digital transfers, static physical storage, or legacy public-key cryptography (RSA/ECC) that is vulnerable both to "
        "human interception today and to Harvest-Now-Decrypt-Later (HNDL) attacks by future Cryptographically Relevant Quantum Computers (CRQCs).\n\n"
        "QuantumShield is an enterprise-grade, post-quantum-cryptography-secured academic paper distribution and anti-leak governance platform. "
        "Engineered using NIST FIPS 203 (ML-KEM-768) for quantum-resistant key encapsulation and NIST FIPS 204 (ML-DSA-65) for non-repudiable digital signatures, "
        "QuantumShield integrates a Zero-Trust Access Policy Engine, a high-security Examination State Machine, and an Immutable SHA3-Linked Event Provenance Chain. "
        "This document presents the project rationale, structural resolution of exam paper leaks, mathematical post-quantum defense, comprehensive capabilities analysis, "
        "and future strategic roadmap."
    )
    r_abs.font.name = "Calibri"
    r_abs.font.size = Pt(11)

    # SECTION 1: NECESSITY & JANTAR MANTAR CONTEXT
    add_h1("1. Necessity & Context: Solving National Exam Paper Leak Crises")
    
    add_h2("1.1 Background: The National Paper Leak Crisis & Jantar Mantar Protests")
    p1 = doc.add_paragraph(
        "The integrity of national competitive examinations (such as NEET-UG, UGC-NET, and State Public Service Commission exams) has repeatedly collapsed "
        "due to widespread paper leaks. Massive student demonstrations at Jantar Mantar, New Delhi, and nationwide academic protests have highlighted four systemic failures "
        "in legacy distribution pipelines:"
    )
    p1.paragraph_format.line_spacing = 1.15

    bullets1 = [
        ("Pre-Exam Storage Interception: ", "Question papers stored on central servers or transferred via insecure channels (email, cloud drives, WhatsApp/Telegram) are accessed days before exam day by compromised staff or external hackers."),
        ("Unregulated Key Distribution: ", "Traditional PDF encryption uses static passwords or shared master keys. Once a single proctor or printing press staff member obtains the key, it is immediately distributed across social media networks."),
        ("Repudiation & Lack of Accountability: ", "System administrators and distribution nodes can delete access logs or deny leaking papers because audit logs are mutable database entries."),
        ("Time Window Violations: ", "Exam centers decrypt and print question papers hours before the designated start time, creating a window during which papers are photographed and solved by coaching mafias.")
    ]
    for b_title, b_text in bullets1:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.line_spacing = 1.15
        r_t = bp.add_run(b_title)
        r_t.bold = True
        bp.add_run(b_text)

    add_h2("1.2 How QuantumShield Resolves Paper Leaks")
    p2 = doc.add_paragraph(
        "QuantumShield restructures paper distribution from a 'trust-by-default' model to a Cryptographic Lockbox with Zero-Trust Gated Key Release:"
    )
    p2.paragraph_format.line_spacing = 1.15

    sol_bullets = [
        ("Ephemeral Key Decapsulation (No Static Passwords): ", "Document Data Encryption Keys (DEKs) are generated as random 256-bit AES keys, encrypted using AES-256-GCM, and encapsulated using ML-KEM-768. Decryption keys NEVER exist static on disk or in database fields; they are decapsulated in RAM only when access policies pass."),
        ("Exam State Machine Lockbox (exams.py): ", "Papers assigned to exams pass through strict state transitions: created -> locked -> released -> closed. Even if an exam centre possesses valid credentials, the backend prohibits decapsulation and download while the exam is in created or locked state. Key release is unlocked ONLY when the exam transitions to released status at the exact scheduled exam time."),
        ("Zero-Trust Access Policy Engine (policies.py): ", "Gates document download against multi-factor rules: User Identity/Role authorization, UTC Time-Window matching (now >= valid_from and now <= valid_until), and strict Download Quotas."),
        ("Immutable Provenance Event Chains (provenance.py): ", "Every lifecycle event (Created, Downloaded, KeyRotated) writes a cryptographic ProvenanceEvent record linking prev_record_hash = SHA3-256(prev_event). Database log tampering breaks chain linkage, flagging TAMPERED_OR_CORRUPT and identifying the exact breach point.")
    ]
    for b_title, b_text in sol_bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.line_spacing = 1.15
        r_t = bp.add_run(b_title)
        r_t.bold = True
        bp.add_run(b_text)

    # SECTION 2: POST-QUANTUM CRYPTOGRAPHY
    add_h1("2. Post-Quantum Cryptography Architecture & HNDL Defense")
    
    add_h2("2.1 The Harvest-Now-Decrypt-Later (HNDL) Threat Model")
    p3 = doc.add_paragraph(
        "State-sponsored adversaries and criminal syndicates are actively capturing encrypted network traffic and database backups containing high-value academic examination materials today. "
        "Under the Harvest-Now-Decrypt-Later (HNDL) strategy, adversaries store intercepted ciphertexts encrypted with legacy public-key algorithms (RSA-2048, ECC P-256). "
        "When a Cryptographically Relevant Quantum Computer (CRQC) executing Shor's Algorithm becomes available, the adversary will factor RSA moduli or solve discrete logarithms in polynomial time, revealing all stored papers."
    )
    p3.paragraph_format.line_spacing = 1.15

    add_h2("2.2 QuantumShield Post-Quantum Security Stack")
    p4 = doc.add_paragraph(
        "QuantumShield neutralizes HNDL threats by replacing vulnerable public-key primitives with NIST-standardized Post-Quantum Cryptography:"
    )
    p4.paragraph_format.line_spacing = 1.15

    pqc_specs = [
        ("NIST FIPS 203 (ML-KEM-768 / Kyber768): ", "Based on the hardness of Module Learning With Errors (M-LWE) over module lattices. Provides Category 3 security (AES-192 equivalent). Quantum algorithms (Grover's) provide only quadratic speedup against symmetric keys, leaving AES-256 and ML-KEM-768 fully secure. Implements Implicit Rejection: on corrupted input, decapsulation deterministically derives a pseudo-random mismatched key, causing downstream AES-GCM tag verification to fail cleanly without side-channel leaks."),
        ("NIST FIPS 204 (ML-DSA-65 / Dilithium3): ", "Based on Module Short Integer Solution (M-SIS) over module lattices using Fiat-Shamir with Aborts. Digitally signs document SHA3-256 digests at rest, proving distributor non-repudiation and guaranteeing existential unforgeability (EUF-CMA)."),
        ("AES-256-GCM & SHA3-256: ", "Authenticated symmetric payload encryption (256-bit key, 96-bit IV, 128-bit tag) paired with FIPS 202 SHA3-256 cryptographic digests.")
    ]
    for b_title, b_text in pqc_specs:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.line_spacing = 1.15
        r_t = bp.add_run(b_title)
        r_t.bold = True
        bp.add_run(b_text)

    # SECTION 3: CAPABILITIES MATRIX (TABLE)
    add_h1("3. Capabilities Analysis: Handled Threat Cases vs Limitations")

    p_tbl_intro = doc.add_paragraph(
        "The following matrix categorizes QuantumShield's security capabilities into Handled Supported Threat Vectors vs Explicit Out-of-Scope Hardware Boundaries:"
    )
    p_tbl_intro.paragraph_format.line_spacing = 1.15

    # Table 1: Handled Capabilities vs Out of Scope
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    hdr_titles = ["Threat / Vector", "Category", "QuantumShield Defense Mechanism", "Status"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B") # Dark slate
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    matrix_data = [
        ("Early Pre-Exam Paper Leakage", "Supported Threat", "Exam State Machine locks key release until released exam time.", "BLOCKED"),
        ("Harvest-Now-Decrypt-Later (HNDL)", "Supported Threat", "ML-KEM-768 key encapsulation immune to Shor's algorithm.", "BLOCKED"),
        ("Unauthorized Role Download", "Supported Threat", "Zero-trust policy gates release by role ID & identity claims.", "BLOCKED"),
        ("Quota Exhaustion Attack", "Supported Threat", "Strict download counter check (max_downloads enforcement).", "BLOCKED"),
        ("Database Audit Log Tampering", "Supported Threat", "SHA3 Provenance hash chain detects prev_hash mismatches.", "DETECTED"),
        ("Ciphertext Payload Tampering", "Supported Threat", "AES-256-GCM tag & SHA3 digest verification failure.", "BLOCKED"),
        ("Digital Signature Forgery", "Supported Threat", "ML-DSA-65 signature verification failure.", "BLOCKED"),
        ("Credential Brute Forcing", "Supported Threat", "Bcrypt hashing & slowapi rate-limiting (10 req/min).", "BLOCKED"),
        ("Physical Monitor Screen Photo", "Out-of-Scope Limit", "Software cryptography terminates once bytes render to screen.", "OUT OF SCOPE"),
        ("Physical RAM Probing / DPA", "Out-of-Scope Limit", "Hardware side-channel probes require hardware enclosures.", "OUT OF SCOPE"),
        ("Volumetric Network DDoS", "Out-of-Scope Limit", "Traffic flooding must be mitigated at edge WAF / Anycast layer.", "OUT OF SCOPE")
    ]

    for row_idx, data in enumerate(matrix_data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_color)
            for p in row_cells[c_idx].paragraphs:
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(9)
                    if c_idx == 3:
                        r.font.bold = True
                        if val == "BLOCKED":
                            r.font.color.rgb = RGBColor(22, 101, 52) # Dark green
                        elif val == "DETECTED":
                            r.font.color.rgb = RGBColor(30, 64, 175) # Blue
                        else:
                            r.font.color.rgb = RGBColor(180, 83, 9) # Amber

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 4: ATTACK SIMULATION LAB
    add_h1("4. Demonstrable Attack Simulation Lab Matrix")
    p5 = doc.add_paragraph(
        "QuantumShield includes a dedicated Attack Simulation Lab (/attack-simulation) enabling security administrators to demonstrate "
        "resilience by testing attacks side-by-side against Secured PQC Targets and Vulnerable Legacy Targets:"
    )
    p5.paragraph_format.line_spacing = 1.15

    sim_data = [
        ("Simulate File Tampering", "Flips a single byte in stored ciphertext payload.", "BLOCKED (SHA3 / GCM Tag Fail)", "SUCCEEDED (Payload Breached)"),
        ("Simulate Invalid Signature", "Attaches forged digital signature to document.", "BLOCKED (ML-DSA Verification Fail)", "SUCCEEDED (Signature Forged)"),
        ("Simulate Wrong AES Key", "Decrypts payload with mismatched symmetric key.", "BLOCKED (AES-256-GCM Auth Fail)", "SUCCEEDED (Key Mismatch Unchecked)"),
        ("Simulate Key Corruption", "Corrupts Kyber encapsulated key ciphertext.", "BLOCKED (ML-KEM Implicit Rejection)", "SUCCEEDED (Decapsulation Leaked)"),
        ("Simulate Unauthorized Access", "Downloads paper without required role claims.", "BLOCKED (Zero-Trust Policy Denial)", "SUCCEEDED (Role Unchecked)"),
        ("Execute Key Rotation", "Triggers emergency post-quantum key re-signing.", "RE-SIGNED & PROVENANCE AUDITED", "RE-SIGNING MANDATORY")
    ]

    table2 = doc.add_table(rows=1, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False

    hdr_cells2 = table2.rows[0].cells
    hdr_titles2 = ["Attack Scenario", "Description", "Secured Target (PQC)", "Vulnerable Target (Legacy)"]
    for i, title in enumerate(hdr_titles2):
        hdr_cells2[i].text = title
        set_cell_background(hdr_cells2[i], "1E293B")
        for paragraph in hdr_cells2[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    for row_idx, data in enumerate(sim_data):
        row_cells2 = table2.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(data):
            row_cells2[c_idx].text = val
            set_cell_background(row_cells2[c_idx], bg_color)
            for p in row_cells2[c_idx].paragraphs:
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.size = Pt(9)
                    if c_idx == 2:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(22, 101, 52)
                    elif c_idx == 3:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(185, 28, 28)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 5: FUTURE STRATEGIC ROADMAP
    add_h1("5. Future Strategic Roadmap")
    p6 = doc.add_paragraph(
        "To advance QuantumShield toward national-scale deployment across competitive examination boards, the following three strategic phases are planned:"
    )
    p6.paragraph_format.line_spacing = 1.15

    roadmap_items = [
        ("Phase 1: Steganographic Dynamic Watermarking: ", "Injecting invisible, centre-specific micro-watermarks (centre ID, timestamp, proctor IP) directly into the decrypted PDF canvas stream before rendering. If a paper is photographed with a camera, automated extraction algorithms immediately identify the originating exam centre and individual monitor."),
        ("Phase 2: FIPS 140-3 Level 3 HSM & TPM 2.0 Integration: ", "Moving server private seed files (server_mldsa_seed.bin and system_mlkem_seed.bin) into dedicated Hardware Security Modules (HSMs) or Trusted Platform Modules (TPMs), ensuring private key bytes cannot be extracted even by root OS users."),
        ("Phase 3: Multi-Party Threshold Cryptography (MPC / k-of-n Secret Sharing): ", "Requiring threshold approvals from k out of n independent exam board trustees to authorize key decapsulation, eliminating single points of administrative failure.")
    ]
    for b_title, b_text in roadmap_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.line_spacing = 1.15
        r_t = bp.add_run(b_title)
        r_t.bold = True
        bp.add_run(b_text)

    # SECTION 6: CONCLUSION
    add_h1("6. Conclusion")
    p_conc = doc.add_paragraph(
        "QuantumShield addresses the systemic vulnerabilities responsible for national academic paper leaks. "
        "By replacing insecure digital distribution pipelines with NIST-standardized Post-Quantum Cryptography (ML-KEM-768 & ML-DSA-65), "
        "State Machine Gated Key Release, Zero-Trust Access Policies, and Immutable SHA3 Provenance Chains, QuantumShield guarantees that "
        "academic papers remain cryptographically sealed until the exact moment of examination, resistant to both human interception today and quantum decryption tomorrow."
    )
    p_conc.paragraph_format.line_spacing = 1.15

    out_file = "c:\\Users\\anubh\\Desktop\\cypto\\QuantumShield_Project_Report.docx"
    doc.save(out_file)
    print(f"Word Document created successfully at: {out_file}")

if __name__ == "__main__":
    create_report_docx()
